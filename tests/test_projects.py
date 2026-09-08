import hashlib
import json
import stat
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from generate_pptx import create_presentation
from storyboard_studio.projects import (
    ProjectPayload,
    materialize_project,
    project_from_files,
    read_project,
    safe_project_path,
    unpack_project,
    write_project,
)
from storyboard_studio.resources import demo_outline_path
from storyboard_studio.story import read_story_or_presentation


def visual_project():
    source = Path("assets/demo/native-visuals.json")
    story, _ = read_story_or_presentation(source)
    return project_from_files(story, source.parent)


def test_portable_assets_regenerate_in_an_unrelated_directory(tmp_path):
    project = visual_project()
    archive = tmp_path / "project.zip"
    write_project(project, archive, render=True)
    reopened = read_project(archive)
    assert reopened.story.presentation.assets == project.story.presentation.assets
    assert reopened.files == project.files
    with zipfile.ZipFile(archive) as packed:
        assert {"deck.pptx", "deck.receipt.json", "project.json"}.issubset(packed.namelist())
        manifest = json.loads(packed.read("project.json"))
        for name, digest in manifest["files"].items():
            assert hashlib.sha256(packed.read(name)).hexdigest() == digest
    story_path = unpack_project(archive, tmp_path / "other é directory")
    story, _ = read_story_or_presentation(story_path)
    output = create_presentation(
        story.presentation.model_dump(), tmp_path / "regenerated.pptx", asset_root=story_path.parent
    )
    deck = Presentation(output)
    assert any(shape.has_chart for slide in deck.slides for shape in slide.shapes)
    assert any(shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes)


def test_project_rejects_missing_and_modified_assets(tmp_path):
    project = visual_project()
    missing = project.model_copy(deep=True)
    missing.files.pop(next(iter(missing.files)))
    with pytest.raises(ValueError, match="Missing"):
        materialize_project(missing, tmp_path)
    altered = project.model_copy(deep=True)
    altered.files[next(iter(altered.files))] = "aGVsbG8="
    with pytest.raises(ValueError, match="checksum"):
        materialize_project(altered, tmp_path)


@pytest.mark.parametrize(
    "path",
    [
        "../private",
        "/absolute",
        "C:/private",
        "folder\\file",
        "a//b",
        "NUL.txt",
        "file. ",
        "./data.csv",
        ".",
        "bad?.png",
    ],
)
def test_portable_paths_reject_platform_specific_escape_routes(path):
    with pytest.raises(ValueError):
        safe_project_path(path)


def test_project_archive_rejects_traversal_links_duplicates_and_bombs(tmp_path):
    for index, (name, kind, content) in enumerate(
        [
            ("../outside", "file", b"x"),
            ("link", "symlink", b"/etc/passwd"),
            ("duplicate", "duplicate", b"x"),
            ("large", "file", b"0" * 12_000_001),
        ]
    ):
        path = tmp_path / f"bad-{index}.zip"
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as output:
            member = zipfile.ZipInfo(name)
            if kind == "symlink":
                member.create_system = 3
                member.external_attr = (stat.S_IFLNK | 0o777) << 16
            output.writestr(member, content, compress_type=zipfile.ZIP_DEFLATED)
            if kind == "duplicate":
                output.writestr(name.upper(), b"other")
        with pytest.raises(ValueError):
            read_project(path)
    assert not (tmp_path.parent / "outside").exists()


def test_project_does_not_overwrite_existing_destinations(tmp_path):
    story, _ = read_story_or_presentation(demo_outline_path())
    project = ProjectPayload(story=story)
    archive = tmp_path / "owned.zip"
    archive.write_bytes(b"user owned")
    with pytest.raises(FileExistsError):
        write_project(project, archive)
    assert archive.read_bytes() == b"user owned"
    good = tmp_path / "good.zip"
    write_project(project, good)
    existing = tmp_path / "existing"
    existing.mkdir()
    with pytest.raises(ValueError, match="existing"):
        unpack_project(good, existing)
    assert list(existing.iterdir()) == []


def test_without_sources_does_not_modify_original_story(tmp_path):
    story, _ = read_story_or_presentation(demo_outline_path())
    assert any(slide.sources for slide in story.presentation.slides)
    project = ProjectPayload(story=story, include_sources=False)
    archive = tmp_path / "without-sources.zip"
    write_project(project, archive)
    reopened = read_project(archive)
    assert all(not slide.sources for slide in reopened.story.presentation.slides)
    assert not reopened.story.decision_brief.evidence
    assert any(slide.sources for slide in story.presentation.slides)


def test_bad_zip_is_a_validation_error(tmp_path):
    path = tmp_path / "bad.zip"
    path.write_bytes(b"not a project")
    with pytest.raises(ValueError, match="project archive"):
        read_project(path)


def test_packing_rejects_symlink_assets_even_with_matching_bytes(tmp_path):
    project = visual_project()
    root = tmp_path / "project"
    root.mkdir()
    story = materialize_project(project, root)
    asset = story.presentation.assets[0]
    source = root / asset.path
    outside = tmp_path / "outside-data"
    source.rename(outside)
    source.symlink_to(outside)
    with pytest.raises(ValueError, match="symlink"):
        project_from_files(story, root)
    assert outside.is_file()
