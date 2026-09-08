"""Prove that every public export surface preserves the authored story."""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from pptx import Presentation

import server
from schemas import PresentationPayload
from storyboard_studio.cli import main as cli_main
from storyboard_studio.projects import ProjectPayload
from storyboard_studio.renderer import create_presentation
from storyboard_studio.story import migrate_presentation_v1
from storyboard_studio.tool_server import ToolServer

ROOT = Path(__file__).resolve().parents[1]


def _semantic_outline() -> dict[str, Any]:
    return json.loads((ROOT / "examples/fixtures/semantic-blocks.json").read_text(encoding="utf-8"))


def _text_fragments(presentation: dict[str, Any]) -> list[str]:
    """Return authored values that the renderer promises to keep observable."""

    fragments: list[str] = []

    def add(value: object) -> None:
        if isinstance(value, str) and value.strip():
            fragments.append(value)

    add(presentation.get("title"))
    add(presentation.get("subtitle"))
    for slide in presentation["slides"]:
        add(slide.get("title"))
        add(slide.get("content"))
        for bullet in slide.get("bullet_points", []):
            add(bullet.get("label"))
            add(bullet.get("title"))
            add(bullet.get("description"))
        block = slide.get("content_block") or {}
        block_type = block.get("type")
        if block_type == "standard":
            for point in block["points"]:
                add(point.get("label"))
                add(point.get("title"))
                add(point.get("description"))
        elif block_type == "comparison":
            for side in block["sides"]:
                add(side.get("title"))
                add(side.get("summary"))
            for criterion in block["criteria"]:
                add(criterion.get("label"))
                add(criterion.get("left"))
                add(criterion.get("right"))
        elif block_type == "decision":
            add(block.get("decision"))
            for option in block["options"]:
                add(option.get("title"))
                add(option.get("description"))
            add(block.get("rationale"))
            add(block.get("owner"))
        elif block_type == "timeline":
            for step in block["steps"]:
                add(step.get("label"))
                add(step.get("title"))
                add(step.get("owner"))
        elif block_type == "metric":
            add(block.get("value"))
            add(block.get("label"))
            add(block.get("context"))
            add(block.get("source"))
        elif block_type == "process":
            for step in block["steps"]:
                add(step.get("title"))
                add(step.get("description"))
        elif block_type == "quote":
            add(block.get("quote"))
            add(block.get("attribution"))
            add(block.get("evidence"))
        elif block_type == "table":
            for column in block["columns"]:
                add(column)
            for row in block["rows"]:
                for cell in row["cells"]:
                    add(cell)
            add(block.get("accessible_summary"))
        elif block_type == "chart":
            add(block.get("title"))
            add(block.get("source_note"))
        elif block_type == "image":
            add(block.get("alt_text"))
            add(block.get("caption"))
        add(slide.get("speaker_notes"))
        for source in slide.get("sources", []):
            for key in ("label", "evidence", "owner", "url", "local_reference", "checked_date", "license"):
                add(source.get(key))
    return fragments


def _extract_pptx_text(content: bytes) -> list[str]:
    deck = Presentation(io.BytesIO(content))
    values: list[str] = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                values.extend(line for line in shape.text.splitlines() if line.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values.extend(cell.text for cell in row.cells if cell.text.strip())
        values.extend(line for line in slide.notes_slide.notes_text_frame.text.splitlines() if line.strip())
    return values


def _assert_authored_content_is_preserved(content: bytes, presentation: dict[str, Any]) -> None:
    rendered = _extract_pptx_text(content)
    previous = -1
    for fragment in _text_fragments(presentation):
        # A renderer may combine adjacent authored fields into one text frame
        # (for example rationale and owner), so the next fragment may share the
        # previous frame while still retaining document order.
        matches = [index for index, value in enumerate(rendered) if index >= previous and fragment in value]
        assert matches, f"Authored fragment was lost from PPTX: {fragment!r}"
        previous = matches[0]


def _export_all_entrypoints(presentation: dict[str, Any], root: Path, monkeypatch) -> dict[str, bytes]:
    story = migrate_presentation_v1(PresentationPayload.model_validate(presentation))
    outputs: dict[str, bytes] = {}

    direct = root / "direct.pptx"
    create_presentation(presentation, direct)
    outputs["renderer"] = direct.read_bytes()

    source = root / "source.json"
    source.write_text(json.dumps(presentation), encoding="utf-8")
    cli_output = root / "cli.pptx"
    assert cli_main(["export", "--input", str(source), "--output", str(cli_output)]) == 0
    outputs["cli"] = cli_output.read_bytes()

    cache = root / "http-exports"
    monkeypatch.setattr(server, "OUTPUT_DIR", cache)
    server._requests.clear()
    with TestClient(server.app, base_url="http://127.0.0.1") as client:
        for route in ("/api/presentations", "/api/v1/presentations"):
            response = client.post(route, json={"presentation": presentation})
            assert response.status_code == 201, response.text
            download = client.get(response.json()["download_url"])
            assert download.status_code == 200
            outputs[route] = download.content

        bundle_response = client.post("/api/v1/bundles", json=story.model_dump(mode="json"))
        assert bundle_response.status_code == 201, bundle_response.text
        bundle_download = client.get(bundle_response.json()["download_url"])
        with zipfile.ZipFile(io.BytesIO(bundle_download.content)) as archive:
            outputs["/api/v1/bundles"] = archive.read("deck.pptx")

        project = ProjectPayload(story=story).model_dump(mode="json")
        for route in ("/api/v1/projects/export", "/api/v1/projects/bundle"):
            response = client.post(route, json=project)
            assert response.status_code == 201, response.text
            download = client.get(response.json()["download_url"])
            assert download.status_code == 200
            if route.endswith("bundle"):
                with zipfile.ZipFile(io.BytesIO(download.content)) as archive:
                    outputs[route] = archive.read("deck.pptx")
            else:
                outputs[route] = download.content

    tools = ToolServer(root, Path("tool-output"))
    rendered = tools.handle(
        {
            "id": "parity",
            "action": "render",
            "arguments": {
                "story": story.model_dump(mode="json"),
                "filename": "tool-export.pptx",
                "acknowledge_review_warnings": True,
            },
        }
    )
    assert rendered["ok"] is True
    outputs["jsonl"] = (root / "tool-output" / "tool-export.pptx").read_bytes()
    return outputs


def test_all_public_export_entrypoints_preserve_content_sources_and_order(tmp_path: Path, monkeypatch):
    fixtures = [_semantic_outline()]
    evidence_fixture = ROOT / "examples/fixtures/evidence-edge-cases.json"
    fixtures.append(json.loads(evidence_fixture.read_text(encoding="utf-8")))

    for index, presentation in enumerate(fixtures):
        root = tmp_path / f"fixture-{index}"
        root.mkdir()
        for _entrypoint, content in _export_all_entrypoints(presentation, root, monkeypatch).items():
            _assert_authored_content_is_preserved(content, presentation)
